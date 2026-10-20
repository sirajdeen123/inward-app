"""Materials Return Register (Excel): one row per request, the SKU lines with the four quantities, and the GRN invoices
not yet raised on the portal. Downloaded from the screen and attached to the 09:30 daily report."""
from __future__ import annotations

import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from inward.rma import services as svc
from inward.rma.models import Grn, Request

FILL = {"RED": PatternFill("solid", fgColor="F8D7DA"), "AMBER": PatternFill("solid", fgColor="FFF1D6"),
        "GREEN": PatternFill("solid", fgColor="E7F4EA")}


def _sheet(ws, title, head, rows, flags=None):
    ws.title = title
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E3E5EA")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, r in enumerate(rows):
        ws.append(r)
        f = flags[i] if flags else ""
        if f in FILL:
            for c in ws[ws.max_row]:
                c.fill = FILL[f]
    for i, h in enumerate(head, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:300]])
        ws.column_dimensions[get_column_letter(i)].width = min(45, max(9, width + 2))
    ws.freeze_panes = "B2"


def _d(v):
    return v.strftime("%d-%m-%Y") if v else ""


def build() -> bytes:
    wb = Workbook()
    rws = svc.rows(Request.objects.all())
    grn_info = {}
    for g in Grn.objects.prefetch_related("lines"):
        for ln in g.lines.all():
            if ln.request_id:
                grn_info.setdefault(ln.request_id, g)
    head = ["Request Id", "Requested", "Dealer code", "Dealer", "Invoice", "Invoice date", "Return type", "Value (Rs)",
            "Portal status", "Pending at", "Stage", "Days", "Days at stage", "Flag", "GRN", "GRN date", "Gate entry",
            "Transport", "LR No", "Freight (Rs)", "RMA No", "RMA date", "CN No", "Credited (Rs)", "IRR TRX",
            "Reconciliation", "Sales Officer", "Closed", "Closed as", "Verified by", "Note"]
    out, flags, sku_rows = [], [], []
    for x in rws:
        r, g = x.r, grn_info.get(x.r.pk)
        rec = r.recon if r.closed_at and r.recon else svc.reconcile(r)
        out.append([r.request_id, _d(r.requested_date), r.dealer_code, r.dealer_name, r.invoice_no, _d(r.invoice_date),
                    r.return_type, float(r.total_value), r.status, r.pending_at, x.stage_label, x.days,
                    x.days_at_stage, x.flag, ", ".join(x.grns), _d(g.grn_date) if g else "",
                    g.gate_entry_no if g else "", g.get_transport_mode_display() if g and g.transport_mode else "",
                    g.lr_no if g else "", float(g.freight) if g and g.freight is not None else None, r.rma_no,
                    _d(r.rma_date), ", ".join(rec["cn"]) or r.cn_no, float(rec["credited_value"]) if rec["cn"] else None,
                    ", ".join(rec["trx"]), svc.RECON_TEXT.get(rec["status"], rec["status"]),
                    f"{x.sales_officer[0]} ({x.sales_officer[2]})" if x.sales_officer else "",
                    _d(timezone.localtime(r.closed_at)) if r.closed_at else "", r.close_kind, r.verified_by, r.close_note])
        flags.append(x.flag)
        for ln in rec["lines"]:
            sku_rows.append([r.request_id, r.dealer_name, ln["sku"], ln["description"], ln["condition"], ln["grn"],
                             ln["portal"], ln["cn"], ln["irr"], ln["status"], float(ln["requested_value"] or 0),
                             float(ln["credited_value"] or 0) if rec["cn"] else None, ln["locator"]])
    _sheet(wb.active, "Returns", head, out, flags)
    _sheet(wb.create_sheet(), "SKU lines",
           ["Request Id", "Dealer", "SKU", "Description", "Condition", "GRN pcs", "Portal qty", "CN pcs", "IRR pcs",
            "Check", "Requested (Rs)", "Credited (Rs)", "Locator"], sku_rows)
    un = svc.unlinked_grns()
    _sheet(wb.create_sheet(), "GRN not on portal",
           ["GRN", "GRN date", "Dealer code", "Dealer", "Invoice", "Invoice date", "SKU lines", "Pieces", "Days", "Flag"],
           [[u["grn"].grn_no, _d(u["grn"].grn_date), u["grn"].dealer_code, u["grn"].dealer_name, u["invoice_no"],
             _d(u["invoice_date"]), u["lines"], float(u["pieces"]), u["days"], u["flag"]] for u in un],
           [u["flag"] for u in un])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
