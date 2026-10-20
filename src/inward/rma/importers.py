"""Readers for the Material Return sources (21-Oct-2026). Nothing is saved here - services.py commits.

- GRN PDF: Oracle "GOODS RETURN NOTE", one page per dealer invoice; a bulk export holds many GRNs, pages with the
  same Receipt No belong to one GRN.
- RMA Portal Report: sheets "RMA Summary" (one row per request) and "RMA Details" (SKU lines).
- RMA CN INV Report: one row per credit-note SKU line; Return Order = the portal RMA No.
- Customer Master (BPIL Customer Master Summary) and TSI Master (territory -> Sales Officer).
Headers are found by name, never by position; every file must be for this depot.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from django.conf import settings

from inward.domain.normalize import clean_code, clean_text, parse_date, to_decimal
from inward.importers.common import ImportFailed, header_key

ZERO = Decimal(0)
_DMY = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$")


def depot() -> str:
    return settings.INWARD["DEPOT_CODE"]


def to_date(value) -> date | None:
    d = parse_date(value)
    if d is None and value not in (None, ""):
        m = _DMY.match(clean_text(value))
        if m:
            try:
                d = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            except ValueError:
                d = None
    return d


def num(value) -> Decimal:
    return to_decimal(value) or ZERO


def code_text(value) -> str:
    """Excel numbers come as 13149 or 13149.0 - codes are kept as text without the '.0'."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean_text(value)


def _sheets(path: Path) -> dict[str, list[list]]:
    head = Path(path).read_bytes()[:8]
    if head.startswith(b"PK\x03\x04"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        finally:
            wb.close()
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        import xlrd
        book = xlrd.open_workbook(path)
        out = {}
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                vals = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    vals.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                                if cell.ctype == xlrd.XL_CELL_DATE else (None if cell.ctype == 0 else cell.value))
                rows.append(vals)
            out[sh.name] = rows
        return out
    from inward.importers.common import read_first_sheet
    return {"Sheet1": read_first_sheet(path).rows}


def _header(rows: list[list], required: list[str], optional: list[str] = (), name: str = "") -> tuple[int, dict]:
    wanted = {header_key(n): n for n in [*required, *optional]}
    for i, row in enumerate(rows[:40]):
        found = {}
        for j, cell in enumerate(row):
            key = header_key(cell)
            if key in wanted and wanted[key] not in found:
                found[wanted[key]] = j
        if all(n in found for n in required):
            return i, found
    raise ImportFailed(f"{name}: header row not found (needs {', '.join(required)}). Is this the right report?")


def _cell(row, cols, name):
    i = cols.get(name)
    return row[i] if i is not None and i < len(row) else None


def _text_in(rows, n=6) -> str:
    return " ".join(clean_text(c) for r in rows[:n] for c in r if c not in (None, ""))

# ---------------------------------------------------------------- GRN PDF


@dataclass
class GrnLineData:
    invoice_no: str
    invoice_date: date | None
    sku: str
    description: str
    inv_nop: Decimal
    ret_nop: Decimal
    drums: Decimal
    cartons: Decimal
    bags: Decimal
    loose: Decimal
    weight: Decimal


@dataclass
class GrnData:
    grn_no: str
    grn_date: date | None
    receipt_time: str = ""
    transporter: str = ""
    vehicle_no: str = ""
    depot: str = ""
    dealer_code: str = ""
    dealer_name: str = ""
    site_code: str = ""
    lines: list[GrnLineData] = field(default_factory=list)
    pages: list[int] = field(default_factory=list)


_RNO = re.compile(r"Receipt\s+No\s*:\s*(RTN/\d{4}/\d+)", re.I)
_RDATE = re.compile(r"Receipt\s+Date\s*:\s*(\d{1,2}/\d{1,2}/\d{4})", re.I)
_RTIME = re.compile(r"Receipt\s+Time\s*:\s*(\d{1,2}:\d{2}(?::\d{2})?)", re.I)
_DEPOT = re.compile(r"Return\s+Depot\s*:\s*([0-9A-Z]+)\s*:", re.I)
_TRANS = re.compile(r"Transporter\s*:\s*(.*?)\s+Receipt\s+No", re.I | re.S)
_VEH = re.compile(r"Vehicle\s+No\.?\s*:\s*(.*?)\s+Receipt\s+Time", re.I | re.S)
_HEAD = re.compile(r"Summary\s+(\d{1,12})\s+(\d{1,2}-[A-Z]{3}-\d{2,4})\s+(\d+)\s*-\s*(.*?)\s*\((\d+)\)\s*Inv\s*Nop",
                   re.I | re.S)
_NUM = r"(-?\d+(?:\.\d+)?)"
_ITEM = re.compile(r"([A-Z0-9]{16})\s*-\s*(.*?)\s+" + r"\s+".join([_NUM] * 7)
                   + r"(?=\s+[A-Z0-9]{16}\s*-|\s+TOTAL\b|\s*$)", re.S)


def parse_grn_page(text: str) -> GrnData | None:
    """One page of the Goods Return Note -> header + the invoice's SKU lines (None if not a GRN page).
    The PDF text is read as one string: line breaks differ from page to page."""
    t = clean_text(text)   # also joins the lines
    m = _RNO.search(t)
    if not m or "GOODS RETURN NOTE" not in t.upper():
        return None
    g = GrnData(grn_no=m.group(1).upper(), grn_date=to_date((_RDATE.search(t) or [None, None])[1]))
    g.receipt_time = (_RTIME.search(t) or [None, ""])[1]
    g.depot = (_DEPOT.search(t) or [None, ""])[1]
    g.transporter = clean_text((_TRANS.search(t) or [None, ""])[1])
    g.vehicle_no = clean_text((_VEH.search(t) or [None, ""])[1])
    h = _HEAD.search(t)
    if not h:
        return g
    inv_no, inv_date = h.group(1), to_date(h.group(2))
    g.dealer_code, g.dealer_name, g.site_code = h.group(3), clean_text(h.group(4)), h.group(5)
    body = t[h.end():]
    body = re.sub(r"^\s*Ret\s*Nop\s+Drum/Bkt\s+Crtn\s+Bags\s+Loose\s+Wt", "", body, flags=re.I)
    end = re.search(r"\bTOTAL\b", body)
    body = body[:end.start() + 5] if end else body
    for im in _ITEM.finditer(body):
        n = [Decimal(v) for v in im.groups()[2:]]
        g.lines.append(GrnLineData(inv_no, inv_date, im.group(1), clean_text(im.group(2)), *n))
    return g


def parse_grn_pdf(data: bytes) -> list[GrnData]:
    """All GRNs in an uploaded PDF (single or bulk); pages with the same Receipt No are merged."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        texts = [p.extract_text() or "" for p in reader.pages]
    except Exception as exc:  # noqa: BLE001 - any unreadable PDF
        raise ImportFailed(f"The PDF could not be read ({exc.__class__.__name__}).") from exc
    out: dict[str, GrnData] = {}
    for n, text in enumerate(texts, start=1):
        g = parse_grn_page(text)
        if g is None:
            continue
        if g.depot and g.depot != depot():
            raise ImportFailed(f"Page {n}: this GRN ({g.grn_no}) is for depot {g.depot}, not Depot {depot()}.")
        if g.grn_no in out:
            out[g.grn_no].lines += g.lines
            out[g.grn_no].pages.append(n)
        else:
            g.pages = [n]
            out[g.grn_no] = g
    if not out:
        raise ImportFailed("No Goods Return Note found in this PDF (expected the Oracle GRN print with a Receipt No "
                           "like RTN/2627/287).")
    return list(out.values())

# ---------------------------------------------------------------- RMA Portal Report


@dataclass
class PortalRequest:
    request_id: str
    values: dict
    lines: list[dict] = field(default_factory=list)


@dataclass
class PortalReport:
    as_on: date | None
    requests: dict[str, PortalRequest]


SUMMARY_REQ = ["Request Id", "Depot", "Dealer", "Invoice No", "Status"]
SUMMARY_OPT = ["Region", "Invoice Date", "Total Value", "Return Type", "Material Receipt Date", "Pending At",
               "Requested By", "Requested Date", "ADMIN Action By", "ADMIN Action Date", "RSM Action By",
               "RSM Action Date", "RCM Action By", "RCM Action Date", "RMA No", "RMA Date", "CN No", "CN Date",
               "CN Value", "Oracle Status", "Oracle Error"]
DETAIL_REQ = ["Request Id", "SKU Code", "Quantity"]
DETAIL_OPT = ["SKU Desc", "Value", "Volume", "Status", "Sub Inventory", "Loactor", "Locator", "RMA No", "RMA Date",
              "Depot"]


def _dealer_code(text: str) -> tuple[str, str]:
    t = clean_text(text)
    m = re.match(r"^(\d+)\s*-\s*(.*)$", t)
    return (m.group(1), clean_text(m.group(2))) if m else ("", t)


def _check_depot(value, where: str) -> None:
    text = clean_text(value)
    if text and not re.search(rf"(^|\D){re.escape(depot())}(\D|$)", text):
        raise ImportFailed(f"{where}: depot '{text}' is not Depot {depot()} - only this depot's report is imported.")


def parse_portal(path: Path) -> PortalReport:
    sheets = _sheets(path)
    summ = next((r for n, r in sheets.items() if "SUMMARY" in n.upper()), None)
    det = next((r for n, r in sheets.items() if "DETAIL" in n.upper()), None)
    if summ is None or det is None:
        raise ImportFailed("RMA Portal Report: the sheets 'RMA Summary' and 'RMA Details' are both needed.")
    as_on = None
    m = re.search(r"Report as on\s*-\s*(\d{1,2}/\d{1,2}/\d{4})", _text_in(summ, 2))
    if m:
        as_on = to_date(m.group(1))
    h, cols = _header(summ, SUMMARY_REQ, SUMMARY_OPT, "RMA Summary")
    out: dict[str, PortalRequest] = {}
    for row in summ[h + 1:]:
        rid = code_text(_cell(row, cols, "Request Id"))
        if not rid:
            continue
        _check_depot(_cell(row, cols, "Depot"), f"RMA Summary request {rid}")
        dealer_code, dealer_name = _dealer_code(_cell(row, cols, "Dealer"))
        v = {
            "region": clean_text(_cell(row, cols, "Region"))[:20], "depot": clean_text(_cell(row, cols, "Depot"))[:40],
            "dealer_code": dealer_code, "dealer_name": dealer_name[:200],
            "invoice_no": code_text(_cell(row, cols, "Invoice No"))[:30],
            "invoice_date": to_date(_cell(row, cols, "Invoice Date")),
            "total_value": num(_cell(row, cols, "Total Value")),
            "return_type": clean_text(_cell(row, cols, "Return Type"))[:60],
            "material_receipt_date": to_date(_cell(row, cols, "Material Receipt Date")),
            "status": clean_text(_cell(row, cols, "Status"))[:60],
            "pending_at": clean_text(_cell(row, cols, "Pending At"))[:300],
            "requested_by": clean_text(_cell(row, cols, "Requested By"))[:120],
            "requested_date": to_date(_cell(row, cols, "Requested Date")),
            "admin_by": clean_text(_cell(row, cols, "ADMIN Action By"))[:120],
            "admin_date": to_date(_cell(row, cols, "ADMIN Action Date")),
            "rsm_by": clean_text(_cell(row, cols, "RSM Action By"))[:120],
            "rsm_date": to_date(_cell(row, cols, "RSM Action Date")),
            "rcm_by": clean_text(_cell(row, cols, "RCM Action By"))[:120],
            "rcm_date": to_date(_cell(row, cols, "RCM Action Date")),
            "rma_no": code_text(_cell(row, cols, "RMA No"))[:30], "rma_date": to_date(_cell(row, cols, "RMA Date")),
            "cn_no": code_text(_cell(row, cols, "CN No"))[:40], "cn_date": to_date(_cell(row, cols, "CN Date")),
            "cn_value": to_decimal(_cell(row, cols, "CN Value")) if "CN Value" in cols else None,
            "oracle_status": clean_text(_cell(row, cols, "Oracle Status"))[:60],
            "oracle_error": clean_text(_cell(row, cols, "Oracle Error"))[:300],
        }
        out[rid] = PortalRequest(rid, v)
    h, cols = _header(det, DETAIL_REQ, DETAIL_OPT, "RMA Details")
    for row in det[h + 1:]:
        rid = code_text(_cell(row, cols, "Request Id"))
        sku = clean_code(_cell(row, cols, "SKU Code"))
        if not rid or not sku:
            continue
        _check_depot(_cell(row, cols, "Depot"), f"RMA Details request {rid}")
        line = {"sku": sku[:20], "description": clean_text(_cell(row, cols, "SKU Desc"))[:200],
                "quantity": num(_cell(row, cols, "Quantity")), "value": num(_cell(row, cols, "Value")),
                "volume": to_decimal(_cell(row, cols, "Volume")), "status": clean_text(_cell(row, cols, "Status"))[:60],
                "sub_inventory": clean_text(_cell(row, cols, "Sub Inventory"))[:30],
                "locator": clean_text(_cell(row, cols, "Loactor") or _cell(row, cols, "Locator"))[:60]}
        if rid in out:
            out[rid].lines.append(line)
            r = out[rid].values
            if not r.get("rma_no"):
                r["rma_no"] = code_text(_cell(row, cols, "RMA No"))[:30]
                r["rma_date"] = to_date(_cell(row, cols, "RMA Date"))
    if not out:
        raise ImportFailed("RMA Portal Report: no request rows found.")
    return PortalReport(as_on, out)

# ---------------------------------------------------------------- RMA CN INV Report


CN_REQ = ["Return Order", "Cn Nunber", "Sku", "CN Nop"]
CN_OPT = ["Cust Code", "Cust Name", "Bill To", "Cn Date", "Inv Number", "Inv Date", "Return Reason", "Pack", "Uom",
          "INV Nop", "Ruturn Ltr", "Ruturn KG", "Cn Rate", "Cn Number", "Return Ltr", "Return KG"]


@dataclass
class CnReport:
    lines: list[dict]
    duplicates: int


def parse_cn(path: Path) -> CnReport:
    rows = next(iter(_sheets(path).values()))
    head = _text_in(rows, 8)
    m = re.search(r"Organization Code\s*::\s*([0-9A-Z]+)", head, re.I)
    if m and m.group(1) != depot():
        raise ImportFailed(f"RMA CN INV Report: organization {m.group(1)}, not Depot {depot()}.")
    h, cols = _header(rows, ["Return Order", "Sku", "CN Nop"], CN_OPT + ["Cn Nunber"], "RMA CN INV Report")
    if "Cn Nunber" not in cols and "Cn Number" not in cols:
        raise ImportFailed("RMA CN INV Report: the CN number column was not found.")
    seen, lines, dup = set(), [], 0
    for row in rows[h + 1:]:
        rma, sku = code_text(_cell(row, cols, "Return Order")), clean_code(_cell(row, cols, "Sku"))
        if not rma or not sku:
            continue
        line = {"rma_no": rma[:30], "cust_code": code_text(_cell(row, cols, "Cust Code"))[:20],
                "cust_name": clean_text(_cell(row, cols, "Cust Name"))[:200],
                "bill_to": code_text(_cell(row, cols, "Bill To"))[:20],
                "cn_no": clean_text(_cell(row, cols, "Cn Nunber") or _cell(row, cols, "Cn Number"))[:40],
                "cn_date": to_date(_cell(row, cols, "Cn Date")),
                "inv_no": clean_text(_cell(row, cols, "Inv Number"))[:40], "inv_date": to_date(_cell(row, cols, "Inv Date")),
                "reason": clean_text(_cell(row, cols, "Return Reason"))[:120], "sku": sku[:20],
                "pack": to_decimal(_cell(row, cols, "Pack")), "uom": clean_text(_cell(row, cols, "Uom"))[:10],
                "cn_nop": num(_cell(row, cols, "CN Nop")), "inv_nop": to_decimal(_cell(row, cols, "INV Nop")),
                "return_ltr": to_decimal(_cell(row, cols, "Ruturn Ltr") or _cell(row, cols, "Return Ltr")),
                "return_kg": to_decimal(_cell(row, cols, "Ruturn KG") or _cell(row, cols, "Return KG")),
                "cn_rate": to_decimal(_cell(row, cols, "Cn Rate"))}
        key = tuple(sorted((k, str(v)) for k, v in line.items()))
        if key in seen:      # the Oracle report repeats a line exactly (tax join) - counted once
            dup += 1
            continue
        seen.add(key)
        lines.append(line)
    if not lines:
        raise ImportFailed("RMA CN INV Report: no credit-note lines found.")
    return CnReport(lines, dup)

# ---------------------------------------------------------------- Customer Master and TSI Master


def territory_no(text) -> str:
    t = clean_text(text)
    return t.rsplit(".", 1)[-1] if "." in t else t


def parse_customer_master(path: Path) -> list[dict]:
    rows = next(iter(_sheets(path).values()))
    m = re.search(r"Location\s*:\s*([0-9A-Z]+)\s*:", _text_in(rows, 4), re.I)
    if m and m.group(1) != depot():
        raise ImportFailed(f"Customer Master: location {m.group(1)}, not Depot {depot()}.")
    h, cols = _header(rows, ["TERRITORY", "ACCOUNT NUMBER", "CUSTOMER NAME", "SITE CODE"],
                      ["BL TYPE", "SITE TYPE", "CUSTOMER TYPE", "CITY"], "Customer Master")
    out = {}
    for row in rows[h + 1:]:
        site = code_text(_cell(row, cols, "SITE CODE"))
        if not site:
            continue
        terr = clean_text(_cell(row, cols, "TERRITORY"))
        out[site] = {"site_code": site[:20], "account_no": code_text(_cell(row, cols, "ACCOUNT NUMBER"))[:20],
                     "name": clean_text(_cell(row, cols, "CUSTOMER NAME"))[:200], "territory": terr[:30],
                     "territory_no": territory_no(terr)[:10], "bl_type": clean_text(_cell(row, cols, "BL TYPE"))[:30],
                     "customer_type": clean_text(_cell(row, cols, "CUSTOMER TYPE"))[:60],
                     "city": clean_text(_cell(row, cols, "CITY"))[:80]}
    if not out:
        raise ImportFailed("Customer Master: no sites found.")
    return list(out.values())


_MAIL = re.compile(r"[^@\s<>,;]+@[^@\s<>,;]+\.[^@\s<>,;]+")


def parse_tsi(path: Path) -> list[dict]:
    rows = next(iter(_sheets(path).values()))
    h, cols = _header(rows, ["Territory", "Sales rep"], ["BL", "Email"], "TSI Master")
    out = {}
    for row in rows[h + 1:]:
        terr = territory_no(code_text(_cell(row, cols, "Territory")))
        if not terr:
            continue
        mail = _MAIL.search(clean_text(_cell(row, cols, "Email")))
        out[terr] = {"territory_no": terr[:10], "sales_officer": clean_text(_cell(row, cols, "Sales rep"))[:120],
                     "business_line": clean_text(_cell(row, cols, "BL"))[:30],
                     "email": mail.group(0).lower() if mail else ""}
    if not out:
        raise ImportFailed("TSI Master: no territories found.")
    return list(out.values())


def stamp(value) -> str:
    return value.strftime("%d-%b-%Y") if isinstance(value, (date, datetime)) else ""
