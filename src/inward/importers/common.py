"""Workbook reading and header location shared by all importers.

Headers are located by normalised names, never fixed positions; a missing
required column fails the whole import (SoT 15).
"""
from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from inward.domain.normalize import clean_text


class ImportFailed(Exception):
    """The file cannot be imported at all; nothing may be committed."""


@dataclass
class RowIssue:
    row_no: int  # 1-based sheet row
    message: str
    ref: str | None = None  # business reference the issue blocks (e.g. Ship ID)


@dataclass
class Sheet:
    rows: list[list]  # raw cell values, first sheet only
    sha256: str
    filename: str


@dataclass
class HeaderMap:
    row_index: int  # 0-based index into Sheet.rows
    columns: dict[str, int] = field(default_factory=dict)

    def get(self, row: list, name: str):
        idx = self.columns.get(name)
        return row[idx] if idx is not None and idx < len(row) else None


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


class _HtmlTableReader(HTMLParser):
    """Oracle BI Publisher '.xls' reports are really HTML: text paragraphs above one table.

    Paragraphs outside tables become one-cell rows (report title, date window), table rows become rows.
    """
    _SKIP = {"style", "script", "title", "head"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows: list[list] = []
        self._row: list | None = None
        self._cell: list[str] | None = None
        self._text: list[str] = []
        self._tables = 0
        self._skip = 0

    def _flush_text(self):
        text = " ".join("".join(self._text).split())
        if text:
            self.rows.append([text])
        self._text = []

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip += 1
        elif tag == "table":
            self._flush_text()
            self._tables += 1
        elif tag == "tr":
            self._row = []
        elif tag in ("td", "th"):
            self._cell = []
        elif tag in ("p", "br", "div") and not self._tables:
            self._flush_text()

    def handle_endtag(self, tag):
        if tag in self._SKIP:
            self._skip = max(0, self._skip - 1)
        elif tag in ("td", "th") and self._cell is not None and self._row is not None:
            text = " ".join("".join(self._cell).split())
            self._row.append(text if text else None)
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if any(c is not None for c in self._row):
                self.rows.append(self._row)
            self._row = None
        elif tag == "table":
            self._tables = max(0, self._tables - 1)
        elif tag in ("p", "div") and not self._tables:
            self._flush_text()

    def handle_data(self, data):
        if self._skip:
            return
        if self._cell is not None:
            self._cell.append(data)
        elif not self._tables:
            self._text.append(data)

    def close(self):
        super().close()
        self._flush_text()


def _read_html(path: Path) -> list[list]:
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig", errors="replace")
    reader = _HtmlTableReader()
    reader.feed(text)
    reader.close()
    return reader.rows


def _read_xlsx(path: Path) -> list[list]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_xls(path: Path) -> list[list]:
    import xlrd

    book = xlrd.open_workbook(path)
    sh = book.sheet_by_index(0)
    rows = []
    for r in range(sh.nrows):
        values = []
        for c in range(sh.ncols):
            cell = sh.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                values.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
            elif cell.ctype == xlrd.XL_CELL_EMPTY:
                values.append(None)
            else:
                values.append(cell.value)
        rows.append(values)
    return rows


def detect_format(path: Path) -> str:
    """Decide by content, not extension: Oracle sends HTML files named .xls."""
    with open(path, "rb") as fh:
        head = fh.read(2048)
    if head.startswith(b"PK\x03\x04"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    probe = head.lstrip(b"\xef\xbb\xbf \t\r\n").lower()
    if probe.startswith((b"<html", b"<!doctype html", b"<table", b"<?xml")) or b"<table" in head.lower():
        return "html"
    raise ImportFailed(f"{path.name} is not an Excel workbook or an Oracle HTML report")


def read_first_sheet(path: str | Path) -> Sheet:
    path = Path(path)
    if path.suffix.lower() not in (".xls", ".xlsx"):
        raise ImportFailed(f"Unsupported file type '{path.suffix}'; expected .xls or .xlsx")
    reader = {"xlsx": _read_xlsx, "xls": _read_xls, "html": _read_html}[detect_format(path)]
    return Sheet(rows=reader(path), sha256=file_sha256(path), filename=path.name)


def header_key(value) -> str:
    return clean_text(value).upper().replace(".", "").replace("_", " ")


def locate_header(sheet: Sheet, required: list[str], optional: list[str] = (), scan_rows: int = 40) -> HeaderMap:
    """Find the first row containing every required column name (case/punctuation-insensitive)."""
    wanted = {header_key(n): n for n in [*required, *optional]}
    for i, row in enumerate(sheet.rows[:scan_rows]):
        found = {}
        for j, cell in enumerate(row):
            key = header_key(cell)
            if key in wanted and wanted[key] not in found:
                found[wanted[key]] = j
        if all(n in found for n in required):
            return HeaderMap(row_index=i, columns=found)
    raise ImportFailed(f"Header row not found in {sheet.filename}; required columns: {', '.join(required)}")


def is_blank_row(row: list) -> bool:
    return all(clean_text(c) == "" for c in row)
